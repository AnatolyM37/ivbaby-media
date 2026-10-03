#!/usr/bin/env python3
"""Таблица продвижения @ivbaby_official: ivbaby-promo.xlsx.

  python3 stats.py build  <папка>                  — создать таблицу заново (первичное заполнение)
  python3 stats.py update <папка> [competitors.json] — дописать строку за сегодня в «Динамика»
                                                       и обновить метрики в «Наши посты»
competitors.json — {"mjolk.ru": 262000, ...} (подписчики конкурентов, если собраны браузером).
"""
import sys, os, json, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

FILE = "ivbaby-promo.xlsx"
MSK = datetime.timezone(datetime.timedelta(hours=3))
MAXR = 400  # строки «Динамика», охватываемые формулами и графиками

# (бренд, аккаунт, подписчики 02.10.2026, публикаций, сегмент, что взять)
COMP = [
    ("Mjölk", "mjolk.ru", 262000, 5407, "Малыши и дети, СПб, свой сайт",
     "Полезные шпаргалки («3 слоя на осень», подбор по погоде), призыв «сохраняйте»"),
    ("Acoola", "acoolakids", 206000, 322, "Дети 2–14, сеть магазинов",
     "Короткие подписи-вопросы к Reels, бэкстейдж съёмок"),
    ("Gulliver", "gulliver_wear", 105000, 5296, "Дети и подростки, премиум",
     "Сторителлинг про коллекцию, темы про подростков"),
    ("O'STIN kids", "ostinkids", 86500, 1862, "Дети, масс-маркет",
     "Готовые образы «с чем носить», школьная тема"),
    ("Rudiks", "rudiks.ru", 81000, 1656, "Малыши и дети, свой сайт",
     "Посты от мам-блогеров (UGC), «отправьте подруге», юмор про родителей"),
    ("PlayToday", "playtoday_official", 46700, 2383, "Дети 0–16, магазины",
     "Гиды «как выбрать…», экспертный тон"),
    ("Ohra Kids", "ohrakids", 19900, 330, "Девочки, WB + Lamoda",
     "Артикулы WB прямо в подписи, розыгрыши с партнёрами"),
    ("Batik", "tdbatik", 17400, 1856, "Дети, Урал, свой сайт",
     "Почти только Reels, family look, короткие подписи"),
    ("Ramelka", "ramelka_baby", 5041, 636, "Дети 0–12, WB, своё производство",
     "Показ производства, артикулы WB хэштегами"),
    ("Oziti kids", "ozitikids", 1100, 89, "Дети 104–146, WB",
     "Голосования «1, 2 или 3?» в комментариях, 3 образа из одной вещи"),
]
IV = ("Ivbaby", "ivbaby_official", 3227, 227)

PHRASES = [
    ("Крючок (1-я строка)", "Похолодало? / Заметно похолодало 🍂", "Mjölk", "Да"),
    ("Крючок (1-я строка)", "Тот самый костюм, который…", "Mjölk, Rudiks", "Да"),
    ("Крючок (1-я строка)", "Как выбрать … и не ошибиться?", "PlayToday, Mjölk", "Да"),
    ("Крючок (1-я строка)", "Какой цвет выберет ваш ребёнок?", "Oziti kids, Rudiks", "Да"),
    ("Крючок (1-я строка)", "Школьные брюки, в которых не мёрзнут", "O'STIN kids", "Да"),
    ("Призыв к действию", "Пишите в комментариях 1, 2 или 3 👇", "Oziti kids, Acoola", "Да"),
    ("Призыв к действию", "Сохраняйте, чтобы не потерять 📌", "Mjölk, Rudiks, Моделька", "Да"),
    ("Призыв к действию", "Отправьте подруге / сыну", "Rudiks", "Да"),
    ("Призыв к действию", "Делитесь своими лайфхаками в комментариях", "Mjölk, Gulliver", "План"),
    ("Позиционирование", "Российский производитель детской одежды", "Rudiks, Ramelka", "Да"),
    ("Позиционирование", "Модная / современная / стильная одежда для детей", "Gulliver, O'STIN kids", "Да"),
    ("Позиционирование", "Для детей. С любовью. На каждый день.", "Batik", "Нет"),
    ("Рубрика", "Шпаргалка: размер, слои по погоде, уход", "Mjölk, PlayToday", "Да"),
    ("Рубрика", "Производство / качество в деталях", "Ramelka, Ivbaby (архив)", "План"),
    ("Рубрика", "Family look / парные образы", "Batik, Mjölk", "План"),
    ("Рубрика", "Отзывы и посты мам-блогеров", "Rudiks, Ramelka", "План"),
    ("Рубрика", "Розыгрыш с партнёром", "Ohra Kids", "План"),
    ("Хэштег (бренд)", "#ivbaby #ивбэби", "—", "Да"),
    ("Хэштег (тема)", "#детскаяодежда", "Моделька, Ivbaby", "Да"),
    ("Хэштег (тема)", "#костюмдлядевочки #костюмдлямальчика", "—", "Да"),
    ("Хэштег (тема)", "#костюмсначесом #теплыйкостюм #флисовыйкостюм", "—", "Да"),
    ("Хэштег (тема)", "#брюкипалаццо #школьнаяформа #одеждадляшколы", "—", "Да"),
    ("Хэштег (маркетплейс)", "#wildberries #ozon + #<артикул WB>", "Ramelka, Oziti kids, Ohra Kids", "Да"),
    ("Поисковая фраза (имя профиля)", "Ivbaby | Детская одежда", "Batik, Ramelka, Oziti kids («Детская одежда | …»)", "16.10"),
]

F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def head(ws, row, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font = Font(name=F, bold=True, color="FFFFFF")
        x.fill = HEAD
        x.alignment = Alignment(wrap_text=True, vertical="center")
        x.border = BOX


def style_all(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.font and c.font.name != F:
                    c.font = Font(name=F, bold=c.font.bold, color=c.font.color, size=c.font.size,
                                  underline=c.font.underline, italic=c.font.italic)


def build(folder):
    wb = Workbook()
    accts = [IV] + [(c[0], c[1], c[2], c[3]) for c in COMP]
    n = len(accts)
    today = datetime.datetime.now(MSK).date()

    # ---------- Динамика ----------
    d = wb.active; d.title = "Динамика"
    head(d, 1, ["Дата"] + ["@" + a[1] for a in accts])
    d.cell(row=2, column=1, value=today).number_format = "DD.MM.YYYY"
    for j, a in enumerate(accts, 2):
        c = d.cell(row=2, column=j, value=a[2]); c.number_format = "#,##0"; c.font = Font(name=F, color="0000FF")
    d.cell(row=2, column=2).comment = Comment("Instagram API, 02.10.2026", "Claude")
    d.cell(row=2, column=3).comment = Comment("Конкуренты: страницы профилей в Instagram, 02.10.2026 "
                                              "(у крупных аккаунтов Instagram округляет до тыс.)", "Claude")
    d.column_dimensions["A"].width = 12
    for j in range(2, n + 2):
        d.column_dimensions[get_column_letter(j)].width = 14
    d.freeze_panes = "B2"

    # ---------- Сводка ----------
    s = wb.create_sheet("Сводка", 0)
    s["A1"] = "Продвижение @ivbaby_official — сводка"; s["A1"].font = Font(name=F, bold=True, size=14)
    s["A2"] = ("Данные берутся из листа «Динамика»: последний и предыдущий замер каждого аккаунта. "
               "Новые строки добавляет автозадание; вручную можно дописать строку с датой и подписчиками "
               "(синие цифры на листе «Динамика»).")
    s["A2"].alignment = Alignment(wrap_text=True); s.merge_cells("A2:G2"); s.row_dimensions[2].height = 45
    head(s, 4, ["Аккаунт", "Подписчики сейчас", "Предыдущий замер", "Прирост", "Прирост, %", "Публикаций (02.10)", "Дата замера"])
    cnt = f"COUNTA(Динамика!$A$2:$A${MAXR})"
    for i, a in enumerate(accts):
        r = 5 + i; col = get_column_letter(2 + i)
        rng = f"Динамика!${col}$2:${col}${MAXR}"
        s.cell(row=r, column=1, value="@" + a[1])
        k1 = f'SUMPRODUCT(LARGE(({rng}<>"")*(ROW({rng})-1),1))'
        k2 = f'SUMPRODUCT(LARGE(({rng}<>"")*(ROW({rng})-1),2))'
        s.cell(row=r, column=2, value=f'=IFERROR(INDEX({rng},{k1}),"")').number_format = "#,##0"
        s.cell(row=r, column=3, value=f'=IFERROR(IF({k2}>0,INDEX({rng},{k2}),""),"")').number_format = "#,##0"
        s.cell(row=r, column=4, value=f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),B{r}-C{r},"")').number_format = "+#,##0;-#,##0;0"
        s.cell(row=r, column=5, value=f'=IF(AND(ISNUMBER(D{r}),C{r}>0),D{r}/C{r},"")').number_format = "0.0%;-0.0%;0.0%"
        s.cell(row=r, column=6, value=a[3]).number_format = "#,##0"
        s.cell(row=r, column=7, value=f'=IFERROR(INDEX(Динамика!$A$2:$A${MAXR},{k1}),"")').number_format = "DD.MM.YYYY"
        for cc in range(1, 8):
            s.cell(row=r, column=cc).border = BOX
    s.cell(row=5, column=1).font = Font(name=F, bold=True)
    for col, w in zip("ABCDEFG", [30, 17, 17, 12, 12, 16, 15]):
        s.column_dimensions[col].width = w
    last = 4 + n
    r0 = last + 2
    s.cell(row=r0, column=1, value="Доля Ivbaby от лидера ниши, %").font = Font(name=F, bold=True)
    s.cell(row=r0, column=2, value=f"=IFERROR(B5/MAX(B6:B{last}),\"\")").number_format = "0.00%"

    s["A{}".format(r0 + 2)] = None
    s.sheet_properties.tabColor = "1F3864"
    # ---------- Конкуренты ----------
    k = wb.create_sheet("Конкуренты")
    head(k, 1, ["№", "Бренд", "Аккаунт", "Ссылка", "Подписчики (последний замер)", "Публикаций (02.10)",
                "Сегмент", "Что взять для Ivbaby"])
    for i, c in enumerate(COMP, 1):
        r = i + 1
        k.cell(row=r, column=1, value=i); k.cell(row=r, column=2, value=c[0]); k.cell(row=r, column=3, value="@" + c[1])
        url = f"https://www.instagram.com/{c[1]}/"
        h = k.cell(row=r, column=4, value=url); h.hyperlink = url; h.font = Font(name=F, color="0563C1", underline="single")
        k.cell(row=r, column=5, value=f"=Сводка!B{5 + i}").number_format = "#,##0"
        k.cell(row=r, column=5).font = Font(name=F, color="008000")
        k.cell(row=r, column=6, value=c[3]).number_format = "#,##0"
        k.cell(row=r, column=7, value=c[4]); k.cell(row=r, column=8, value=c[5])
        for cc in range(1, 9):
            k.cell(row=r, column=cc).border = BOX
            k.cell(row=r, column=cc).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGH", [4, 13, 20, 36, 16, 13, 28, 60]):
        k.column_dimensions[col].width = w
    k.cell(row=len(COMP) + 3, column=1,
           value="Источник: открытые профили Instagram, просмотр 02.10.2026. Хэштеги: 7 из 10 брендов не используют их совсем, "
                 "остальные — 0–5 на пост или артикулы WB.").font = Font(name=F, italic=True)
    k.freeze_panes = "C2"

    # ---------- Ключевые фразы ----------
    p = wb.create_sheet("Ключевые фразы")
    head(p, 1, ["Тип", "Фраза / хэштег", "Где встречается", "Используем в Ivbaby"])
    for i, row in enumerate(PHRASES, 2):
        for j, v in enumerate(row, 1):
            c = p.cell(row=i, column=j, value=v); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", [26, 50, 40, 18]):
        p.column_dimensions[col].width = w
    p.freeze_panes = "A2"

    # ---------- Наши посты ----------
    np_ = wb.create_sheet("Наши посты")
    head(np_, 1, ["Дата", "ID в плане", "Товар", "Арт. WB", "Арт. Ozon", "Ссылка", "Охват", "Просмотры",
                  "Лайки", "Комментарии", "Сохранения", "Репосты", "Вовлечённость, %", "Обновлено"])
    for col, w in zip("ABCDEFGHIJKLMN", [11, 9, 38, 12, 12, 34, 9, 11, 8, 12, 12, 9, 15, 12]):
        np_.column_dimensions[col].width = w
    np_.freeze_panes = "C2"
    style_all(wb)
    add_charts(wb)
    path = os.path.join(folder, FILE)
    wb.save(path)
    return path


def add_charts(wb):
    """openpyxl не сохраняет графики при повторном открытии — пересоздаём их по фактическому числу строк."""
    s, d, np_ = wb["Сводка"], wb["Динамика"], wb["Наши посты"]
    s._charts = []
    nd = max(r for r in range(1, MAXR + 1) if r == 1 or d.cell(row=r, column=1).value is not None)
    npost = max(r for r in range(1, np_.max_row + 1) if r == 1 or np_.cell(row=r, column=2).value is not None)
    last = 5 + d.max_column - 2
    lc = LineChart(); lc.title = "Подписчики @ivbaby_official"; lc.y_axis.title = "Подписчики"
    lc.add_data(Reference(d, min_col=2, min_row=1, max_row=max(nd, 2)), titles_from_data=True)
    lc.set_categories(Reference(d, min_col=1, min_row=2, max_row=max(nd, 2)))
    sr = lc.series[0]; sr.smooth = False
    sr.graphicalProperties.line.solidFill = "1F3864"; sr.graphicalProperties.line.width = 28575
    sr.marker.symbol = "circle"; sr.marker.size = 7
    sr.marker.graphicalProperties.solidFill = "1F3864"; sr.marker.graphicalProperties.line.solidFill = "1F3864"
    lc.x_axis.number_format = "DD.MM.YY"; lc.height = 8; lc.width = 18; lc.legend = None
    lc.x_axis.delete = False; lc.y_axis.delete = False
    s.add_chart(lc, f"A{last + 4}")

    bc = BarChart(); bc.type = "bar"; bc.title = "Подписчики: Ivbaby и конкуренты"
    bc.add_data(Reference(s, min_col=2, min_row=4, max_row=last), titles_from_data=True)
    bc.set_categories(Reference(s, min_col=1, min_row=5, max_row=last))
    bc.height = 9; bc.width = 18; bc.legend = None; bc.x_axis.delete = False; bc.y_axis.delete = False
    s.add_chart(bc, f"A{last + 21}")

    rc = BarChart(); rc.title = "Охват наших постов"; rc.y_axis.title = "Аккаунтов"
    rc.add_data(Reference(np_, min_col=7, min_row=1, max_row=max(npost, 2)), titles_from_data=True)
    rc.set_categories(Reference(np_, min_col=2, min_row=2, max_row=max(npost, 2)))
    rc.height = 8; rc.width = 16; rc.legend = None; rc.x_axis.delete = False; rc.y_axis.delete = False
    s.add_chart(rc, "I4")


def _api(folder):
    sys.path.insert(0, folder)
    from publish import load_env, call, API
    return load_env(os.path.join(folder, "api-keys.txt")), call, API


def update(folder, comp_json=None):
    path = os.path.join(folder, FILE)
    wb = load_workbook(path)
    env, call, API = _api(folder)
    tok = env["IG_ACCESS_TOKEN"]
    me = call("GET", API + "/me", {"fields": "followers_count", "access_token": tok})
    comp = json.load(open(comp_json)) if comp_json and os.path.exists(comp_json) else {}

    d = wb["Динамика"]
    heads = [d.cell(row=1, column=j).value for j in range(1, d.max_column + 1)]
    today = datetime.datetime.now(MSK).date()
    r = 2
    while d.cell(row=r, column=1).value is not None:
        v = d.cell(row=r, column=1).value
        if hasattr(v, "date") and v.date() == today or v == today:
            break
        r += 1
    d.cell(row=r, column=1, value=today).number_format = "DD.MM.YYYY"
    for j, h in enumerate(heads[1:], 2):
        acc = h.lstrip("@")
        val = me["followers_count"] if acc == "ivbaby_official" else comp.get(acc)
        if val is not None:
            c = d.cell(row=r, column=j, value=int(val)); c.number_format = "#,##0"
            c.font = Font(name=F, color="0000FF")

    # Наши посты
    plan = json.load(open(os.path.join(folder, "plan.json"), encoding="utf-8"))
    ws = wb["Наши посты"]
    rows = {ws.cell(row=i, column=2).value: i for i in range(2, ws.max_row + 1) if ws.cell(row=i, column=2).value}
    for p in plan["posts"]:
        if p.get("status") != "published" or not p.get("media_id"):
            continue
        i = rows.get(p["id"]) or (max(rows.values(), default=1) + 1)
        rows[p["id"]] = i
        m = call("GET", f"{API}/{p['media_id']}", {"fields": "like_count,comments_count", "access_token": tok})
        try:
            ins = call("GET", f"{API}/{p['media_id']}/insights", {"metric": "reach,views,saved,shares", "access_token": tok})
            ins = {x["name"]: x["values"][0]["value"] for x in ins["data"]}
        except Exception:
            ins = {}
        vals = [datetime.date.fromisoformat(p["date"]), p["id"], p.get("title"), p.get("nmId"), p.get("ozonSku"),
                p.get("permalink"), ins.get("reach"), ins.get("views"), m.get("like_count"), m.get("comments_count"),
                ins.get("saved"), ins.get("shares")]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=int(v) if j == 5 and v else v)
            c.border = BOX; c.font = Font(name=F)
        ws.cell(row=i, column=1).number_format = "DD.MM.YYYY"
        if p.get("permalink"):
            ws.cell(row=i, column=6).hyperlink = p["permalink"]
            ws.cell(row=i, column=6).font = Font(name=F, color="0563C1", underline="single")
        ws.cell(row=i, column=13, value=f'=IF(N(G{i})>0,(I{i}+J{i}+K{i}+L{i})/G{i},"")').number_format = "0.0%"
        ws.cell(row=i, column=14, value=today).number_format = "DD.MM.YYYY"
        for j in (13, 14):
            ws.cell(row=i, column=j).border = BOX; ws.cell(row=i, column=j).font = Font(name=F)
    add_charts(wb)
    wb.save(path)
    return path


if __name__ == "__main__":
    cmd, folder = sys.argv[1], sys.argv[2]
    if cmd == "build":
        print(build(folder))
    else:
        print(update(folder, sys.argv[3] if len(sys.argv) > 3 else None))
